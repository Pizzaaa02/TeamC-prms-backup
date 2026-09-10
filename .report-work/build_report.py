from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

REFERENCE = r"C:\Users\Chitty\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.1\skills\artifact-template-design-report\assets\reference.docx"
OUTPUT = r"C:\Users\Chitty\source\repos\TeamC-prms\PRMS_Project_Part_B_Implementation_Report.docx"

doc = Document(REFERENCE)
p = doc.paragraphs

replacements = {
    2: "PRMS Project Part B Implementation Report",
    4: "Contents",
    5: "Executive summary\nImplementation scope\nCompleted system capabilities\nValidation results\nCompliance assessment\nRemaining production work\nAppendix",
    6: "Executive summary",
    7: "The Property Rental Management System has progressed from the planning and design work documented in Project Part A into a functioning full stack implementation for Project Part B. The current solution provides role-based experiences for Administrators, Landlords, Tenants, and Agents, supported by a React frontend, an Express and TypeScript backend, Prisma data access, and a SQLite development database.",
    8: "The implemented system now covers property discovery and management, booking approval, rental agreements, payment and invoice records, maintenance requests, messaging, notifications, reporting, identity review, and personal-data controls. Public and authenticated pages were tested through the running website, and the final automated validation completed with successful frontend and backend builds, 66 backend tests, 9 frontend authentication tests, and a clean ESLint error check.",
    9: "At a glance",
    10: "Four role-specific portals are connected to live backend data and enforce role and ownership boundaries.",
    11: "Sixty realistic Malaysian property records were added across six required categories, with MYR pricing and map coordinates.",
    12: "PDPA-supporting controls now include notice, consent history, data export, data-subject requests, incident tracking, retention cleanup, and privacy administration.",
    13: "Implementation scope",
    14: "The work reviewed the Sprint 1 to Sprint 3 reports from Project Part A and compared their planned modules with the existing repository. Missing routes, inconsistent data fields, incomplete role pages, authorization gaps, unresponsive layouts, and disconnected demonstration content were identified and corrected in staged implementation passes.",
    15: "All work remains local in the project workspace. No commit or push has been performed, in accordance with the project owner's instruction. The development servers were also fully stopped after verification and can be restarted when work resumes.",
    16: "Completed system capabilities",
    17: "The implementation delivers the following working capabilities across the application.",
    18: "Frontend and role experiences",
    19: "Responsive public pages, registration, authentication, navigation, property browsing, and role-specific dashboards are available for Admin, Landlord, Tenant, and Agent users. The Tenant and Agent dashboards were converted from outdated demonstration cards to live bookings, properties, payments, favourites, and maintenance data. Mobile horizontal overflow was corrected on property browsing and registration pages.",
    20: "Backend workflows and access control",
    21: "Booking confirmation now creates the related rental agreement, pending payment, invoice, and notifications. Payment, booking, reporting, maintenance, messaging, favourite, and notification access was tightened so users receive only authorised records. Agent category access and assigned-property workflows were connected, and landlord reporting is restricted to owned properties.",
    23: "Key takeaway. The main Sprint 1 to Sprint 3 functional scope is now represented by working full stack flows rather than isolated interface mock-ups.",
    24: "",
    25: "Database and Malaysian data",
    26: "The Prisma schema now includes rental agreements, privacy consent, data-subject requests, data-breach incidents, KYC review fields, property coordinates, and maintenance relationships. The development database contains 60 newly seeded Malaysian listings across Terrace, Double Storey, Apartment, Condominium, Semi-D, and Bungalow categories, plus two preserved pre-existing records. System currency was changed to MYR.",
    27: "Validation results",
    28: "Validation was performed at code, API, and browser levels. Public pages and the four authenticated portals were exercised against the running local backend. The results below reflect the final verification completed after the dashboard and authorization corrections.",
    29: "Backend verification. TypeScript compilation completed successfully and all 66 backend tests passed across three suites.",
    30: "Frontend verification. The Vite production build completed successfully, all 9 authentication tests passed, and ESLint reported no errors.",
    31: "Browser verification. Admin, Landlord, Tenant, and Agent dashboards and their principal workflow pages loaded successfully. CORS, public settings, responsive overflow, landlord dashboard loading, tenant payment mapping, and Agent category authorization defects found during testing were corrected.",
    32: "Compliance assessment",
    33: "The application provides technical controls that support Malaysia's Personal Data Protection Act principles: a visible privacy notice, recorded consent choices, access and correction request handling, consent withdrawal and objection requests, personal-data export, retention cleanup, incident registration, and administrative request tracking. Google Maps is click-to-load and informs the user that third-party data transfer may occur.",
    34: "Technical controls alone do not establish complete legal compliance. Before production use, the organisation must confirm its lawful purposes, privacy contact or Data Protection Officer arrangements, retention schedule, processor agreements, security operations, breach-notification procedure, hosting location, and final legal wording.",
    35: "Remaining production work",
    36: "The core implementation is complete for continued development and academic demonstration. The following items remain deployment and operational tasks rather than missing local application modules.",
    37: "Configure a production database, HTTPS hosting, environment secrets, backups, monitoring, email delivery, and a real payment provider.",
    38: "Replace placeholder organisational privacy contacts and policy wording with approved details, complete final user acceptance testing, and optimise the main frontend bundle through additional code splitting if required.",
    39: "Appendix",
    40: "Technology stack: React 19, Vite, Axios, Express, TypeScript, Prisma, SQLite development database, Jest, and ESLint.",
    41: "Final automated result: backend build passed; 66 of 66 backend tests passed; frontend build passed; 9 of 9 frontend authentication tests passed; frontend lint passed.",
    42: "Runtime status at report preparation: frontend and backend development servers are stopped. Source changes and the development database remain uncommitted and unpushed.",
    43: "Reference basis: COS40005 Sprint 1 Deliverables, Group 14 Sprint Report 2, PRMS Sprint 3 Report, the PRMS repository implementation, and official Malaysian Personal Data Protection guidance reviewed during development.",
    45: "Prepared for the continuation of COS40005 Project Part B implementation and review.",
}

for idx, text in replacements.items():
    p[idx].text = text

p[24].style = doc.styles['normal']

for idx in (37, 38):
    p[idx].text = f"• {p[idx].text}"

# Remove the template's generated TOC, which otherwise retains its original
# placeholder headings. The concise contents list above remains editable.
for sdt in list(doc.element.body.xpath('./w:sdt')):
    sdt.getparent().remove(sdt)

# Remove the template's decorative stock-image front page and its section
# break so the project-specific title page becomes the cover.
for paragraph in (p[0], p[1]):
    paragraph._element.getparent().remove(paragraph._element)

# Subtitle and preparation metadata.
t0 = doc.tables[0]
t0.cell(0, 0).text = "Full stack implementation completion and validation report"
t0.cell(0, 2).text = "Prepared by Group 14\nSeptember 2026"

# Findings matrix.
t1 = doc.tables[1]
matrix = [
    ["Area", "Completed implementation", "Current outcome"],
    ["Core workflows", "Properties, bookings, agreements, payments, maintenance, messaging, notifications, reports", "Role-specific full stack flows operate with live records"],
    ["Security and privacy", "Authentication, RBAC, ownership filtering, KYC review, consent and privacy operations", "Access is restricted by role and record ownership"],
    ["Quality", "Responsive fixes, Malaysian seed data, browser testing, builds, tests and lint", "Application is ready for the next academic review stage"],
]
for row, values in zip(t1.rows, matrix):
    for cell, value in zip(row.cells, values):
        cell.text = value
        cell.vertical_alignment = 1

# Start the findings matrix on a fresh page so no table row is split across
# pages. The template repeats the table header automatically.
p[26].add_run().add_break(WD_BREAK.PAGE)

# Replace the template's running-header placeholders.
for section in doc.sections:
    for table in section.header.tables:
        if table.rows and len(table.rows[0].cells) >= 2:
            table.cell(0, 0).text = "PRMS Project Part B Implementation Report"
            table.cell(0, -1).text = "September 2026"
    for text_node in section.header._element.xpath('.//w:t'):
        if text_node.text == 'Report title':
            text_node.text = 'PRMS Project Part B Implementation Report'
        elif text_node.text == 'Date':
            text_node.text = 'September 2026'

# Enforce black headings and professional typography while preserving template styles.
for style_name in ('Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[style_name]
    style.font.color.rgb = RGBColor(0, 0, 0)
for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        run.font.name = 'Helvetica Neue'
        run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Helvetica Neue')
        run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Helvetica Neue')

doc.core_properties.title = "PRMS Project Part B Implementation Report"
doc.core_properties.subject = "Completed full stack implementation and validation"
doc.core_properties.author = "Group 14"
doc.core_properties.comments = ""
doc.save(OUTPUT)
print(OUTPUT)
